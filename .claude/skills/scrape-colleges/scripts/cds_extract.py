"""Pull just the sections Uniseek needs out of a Common Data Set PDF, and render
those pages as images.

    python .claude/skills/scrape-colleges/scripts/cds_extract.py <url-or-path> [--render 10,16] [--no-images]

Downloads the PDF (cached in data/colleges/.cache/cds/), extracts each page's text with
`pdftotext -layout` (falls back to pypdf), and prints each target section tagged with the
PDF PAGE it sits on. A full CDS is 40+ pages; this keeps it to the ~8 that matter.

It also renders those pages to PNG (pypdfium2, falling back to PyMuPDF) and prints the
paths. READ THE IMAGES for every table: text extraction scrambles checkbox grids (C7, C8A,
E1, C21/C22) and multi-column tables (C9, H2/H2A, I3), and some PDFs use fonts that come
out as gibberish. The text dump is for finding things; the images are for reading values.
--render N,M renders extra pages (e.g. when a section spills onto the next page).
"""

import hashlib
import os
import re
import shutil
import subprocess
import sys
import urllib.request

CACHE = os.path.join("data", "colleges", ".cache", "cds")


def h(code: str) -> str:
    """Heading regex for a CDS item: matches 'C7 ', 'C7.', 'C7:', 'I-3.' at line start."""
    return rf"^\s*{code}[.:\s]"


# (label, start pattern, end pattern). Several start alternatives, because templates
# vary ("C7." vs "C7 ") and some PDFs garble headings beyond recognition.
SECTIONS = [
    ("C1 applicants / admitted (exact admit rate; residency table)",
     h("C1") + r"|first-year (students|men|males|women|females)[^\n]{0,40} who applied", h("C2")),
    ("C7 admission factors", h("C7"), h("C8") + r"|C8:"),
    ("C8 test policy", r"^\s*C8:|" + h("C8"), h("C9") + r"|C9-C12"),
    ("C9 test score percentiles", h("C9"), r"Percent of first-time, first-year students with scores in each range|" + h("C10")),
    ("C13-C20 application fee / closing dates / rolling", h("C13") + r"|C13-C20", h("C21") + r"|C21-C22"),
    ("C21-C22 early decision / early action", h("C21") + r"|C21-C22", r"D\. TRANSFER ADMISSION|" + h("D1")),
    ("E1 special study options (study abroad)", h("E1"), h("F1") + r"|F\. STUDENT LIFE"),
    ("F1 student life (housing, Greek)", h("F1"), h("F2") + "|" + h("F3") + r"|ROTC"),
    ("H2 row A: degree-seeking first-years (merit-aid denominator)", h("H2") + r"\s*Number of Enrolled", h("H2A")),
    ("H2A non-need-based aid (row N)", h("H2A"), h("H3") + "|" + h("H4") + r"|^Note: These are the graduates"),
    ("I3 undergraduate class size", h("I-?3") + r"|Undergraduate Class Size", h("I-?4") + r"|J\. DISCIPLINES|" + h("J1")),
]
MAX_LINES = 90


def kind(data: bytes) -> str | None:
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:2] == b"PK":  # zip container → .xlsx
        return "xlsx"
    return None


def fetch(src: str) -> tuple[str, str]:
    """Returns (local path, 'pdf' | 'xlsx'). Checks the bytes BEFORE caching under an extension."""
    if os.path.exists(src):
        with open(src, "rb") as f:
            k = kind(f.read(8))
        if not k:
            sys.exit(f"Not a PDF or xlsx: {src}")
        return src, k
    os.makedirs(CACHE, exist_ok=True)
    stem = os.path.join(CACHE, hashlib.sha1(src.encode()).hexdigest()[:16])
    for k in ("pdf", "xlsx"):
        if os.path.exists(f"{stem}.{k}"):
            return f"{stem}.{k}", k
    ua = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/129.0 Safari/537.36"
    try:
        req = urllib.request.Request(src, headers={"User-Agent": ua})
        with urllib.request.urlopen(req, timeout=60) as r:
            data = r.read()
    except Exception as e:  # SSL-chain or 403 quirks: curl often succeeds where urllib doesn't
        if not shutil.which("curl"):
            raise
        r = subprocess.run(["curl", "-sSL", "--max-time", "90", "-A", ua, src], capture_output=True)
        data = r.stdout
        if r.returncode != 0 or not data:
            sys.exit(f"Download failed ({e}); curl also failed: {r.stderr.decode(errors='replace')[:200]}")
    k = kind(data)
    if not k:
        sys.exit(f"Not a PDF or xlsx: {src}\n(If it's an HTML CDS page, read it with WebFetch.)")
    with open(f"{stem}.{k}", "wb") as f:
        f.write(data)
    return f"{stem}.{k}", k


def xlsx_lines(path: str) -> list[tuple[str, str, str]]:
    """Each non-empty row as ('Sheet!row', values joined (for matching), 'A1=text  C1=X' (for display)).
    Hidden columns are included: some CDS workbooks keep a machine-readable answer key there."""
    import openpyxl  # pip install openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            cells = [c for c in row
                     if getattr(c, "coordinate", None) and c.value is not None and str(c.value).strip()]
            if cells:
                out.append((f"{ws.title}!{cells[0].row}",
                            "  ".join(str(c.value).strip() for c in cells),
                            "  ".join(f"{c.coordinate}={str(c.value).strip()}" for c in cells)))
    return out


def page_count(pdf: str) -> int:
    try:
        import pypdfium2
        return len(pypdfium2.PdfDocument(pdf))
    except ImportError:
        from pypdf import PdfReader
        return len(PdfReader(pdf).pages)


def pages(pdf: str) -> list[str]:
    # One pdftotext call per page: splitting a whole-document dump on form feeds is
    # unreliable (some PDFs emit extra \f), which shifts every page number after it.
    n = page_count(pdf)
    if shutil.which("pdftotext"):
        out = []
        for i in range(1, n + 1):
            r = subprocess.run(["pdftotext", "-layout", "-f", str(i), "-l", str(i), pdf, "-"],
                               capture_output=True, text=True, encoding="utf-8", errors="replace")
            out.append(r.stdout if r.returncode == 0 else "")
        if any(p.strip() for p in out):
            return out
    from pypdf import PdfReader  # pip install pypdf
    return [p.extract_text(extraction_mode="layout") or "" for p in PdfReader(pdf).pages]


def render(pdf: str, page_nos: list[int]) -> list[str]:
    outdir = os.path.splitext(pdf)[0]
    os.makedirs(outdir, exist_ok=True)
    paths = []
    try:
        import pypdfium2
        doc = pypdfium2.PdfDocument(pdf)
        doc.init_forms()  # draw fillable-form values (checkbox marks) — some CDSs are fillable PDFs
        for n in page_nos:
            path = os.path.join(outdir, f"p{n}.png")
            if not os.path.exists(path):
                doc[n - 1].render(scale=1.6).to_pil().save(path)
            paths.append(path)
    except ImportError:
        import fitz  # PyMuPDF
        doc = fitz.open(pdf)
        for n in page_nos:
            path = os.path.join(outdir, f"p{n}.png")
            if not os.path.exists(path):
                doc[n - 1].get_pixmap(dpi=115).save(path)
            paths.append(path)
    return paths


def main():
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    args = sys.argv[1:]
    if not args or args[0].startswith("--"):
        sys.exit(__doc__)
    pdf, fmt = fetch(args[0])
    extra = []
    if "--render" in args:
        extra = [int(x) for x in args[args.index("--render") + 1].split(",") if x.strip()]

    display = {}
    if fmt == "xlsx":
        rows = xlsx_lines(pdf)
        lines = [(loc, text) for loc, text, _ in rows]  # (sheet!row, values)
        display = {loc: disp for loc, _, disp in rows}
        print(f"# CDS extract (xlsx) — {args[0]}\n# local copy: {pdf}\n"
              "# Rows are CELL=value. Read values straight from the cells (no images needed), but confirm\n"
              "# which column is which from each table's header row before mapping an X to a rating.\n")
    else:
        page_texts = pages(pdf)
        lines = []  # (page_no, text)
        for i, page in enumerate(page_texts, start=1):
            lines.extend((i, ln.rstrip()) for ln in page.splitlines())
        print(f"# CDS extract — {args[0]}\n# local copy: {pdf}  ({len(page_texts)} pages)\n")
    head = " ".join(t for _, t in lines[:60])
    m = re.search(r"Common Data Set\s*(20\d\d)\s*[-–]\s*(20)?(\d\d)", head)
    if m:
        print(f"# CDS year detected: {m.group(1)}-{m.group(3)}\n")

    wanted = set(extra)
    for label, start, end in SECTIONS:
        idx = next((i for i, (_, t) in enumerate(lines) if re.search(start, t)), None)
        if idx is None:
            print(f"## {label}: NOT FOUND in text — find it in the page images (headings may be garbled)\n")
            continue
        block = []
        for j in range(idx, min(idx + MAX_LINES, len(lines))):
            if j > idx + 1 and re.search(end, lines[j][1]):
                break
            if lines[j][1].strip():
                block.append(lines[j])
        if fmt == "xlsx":
            print(f"## {label}  ({block[0][0]} to {block[-1][0]})")
            print("\n".join(f"{loc}: {display[loc]}" for loc, _ in block))
            print()
            continue
        pgs = sorted({p for p, _ in block})
        wanted.update(pgs)
        print(f"## {label}  (PDF page{'s' if len(pgs) > 1 else ''} {', '.join(map(str, pgs))})")
        print("\n".join(t for _, t in block))
        print()

    if fmt == "pdf" and "--no-images" not in args and wanted:
        try:
            paths = render(pdf, sorted(wanted))
            print("# Page images — open with the Read tool and read every table from these:")
            for p in paths:
                print(f"#   {p}")
        except ImportError:
            print("# Could not render pages: `pip install pypdfium2` (or pymupdf), then re-run.")


if __name__ == "__main__":
    main()
